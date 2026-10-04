"""Build the concise v0.5.1 core guide from the reviewed Markdown source."""
from pathlib import Path
import re,json
from PIL import Image,ImageDraw,ImageFont
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
ASSETS=ROOT/'docs/product-assets';ASSETS.mkdir(exist_ok=True)
OUT=ROOT/'output/docx/Circular_Intelligence_Core_Guide_v051.docx';OUT.parent.mkdir(parents=True,exist_ok=True)
FONT=Path('C:/Windows/Fonts/arial.ttf');BOLD=Path('C:/Windows/Fonts/arialbd.ttf')

def diagram(name,size,nodes,arrows):
 im=Image.new('RGB',size,'white');d=ImageDraw.Draw(im)
 for a,b in arrows:
  d.line([a,b],fill='#566c65',width=4)
  x,y=b;dx,dy=b[0]-a[0],b[1]-a[1];length=(dx*dx+dy*dy)**.5;ux,uy=dx/length,dy/length
  d.polygon([(x,y),(x-ux*14-uy*7,y-uy*14+ux*7),(x-ux*14+uy*7,y-uy*14-ux*7)],fill='#566c65')
 for (x,y,w,h,title,body) in nodes:
  d.rounded_rectangle((x,y,x+w,y+h),radius=14,fill='#f2f5f0',outline='#83978e',width=2)
  for yy,text,font in [(y+17,title,ImageFont.truetype(str(BOLD),27)),(y+58,body,ImageFont.truetype(str(FONT),23))]:
   lines=text.split('\n')
   for i,line in enumerate(lines):
    tw=d.textlength(line,font=font);d.text((x+(w-tw)/2,yy+i*30),line,font=font,fill='black')
 im.save(ASSETS/name)

diagram('import-flow.png',(1600,520),[
 (20,30,350,150,'Source pack','Inventories and Markdown\nOriginal inputs preserved'),
 (620,30,350,150,'Audit and identity','URLs and hashes\nListings and memberships'),
 (1220,30,350,150,'Candidate corpus','Pages and quality flags\nReviews and provenance'),
 (1220,320,350,150,'Build indexes','Chunks and embeddings\nChroma and SQLite'),
 (620,320,350,150,'Validate','Integrity and retrieval\nLive checks and source gaps'),
 (20,320,350,150,'Activate','Switch configuration\nRestart and retain rollback')
], [((370,105),(620,105)),((970,105),(1220,105)),((1395,180),(1395,320)),((1220,395),(970,395)),((620,395),(370,395))])
diagram('chunk-flow.png',(1600,450),[
 (20,35,380,145,'Circular A','Page 1 and page 2\nOne canonical document ID'),
 (600,35,430,145,'Attributed chunks','Page 1 chunks A1 and A2\nPage 2 chunks A3 and A4'),
 (1210,35,350,145,'Citations','Circular A and source page\nWord spans and neighbours'),
 (20,270,380,145,'Circular B','Separate document ID\nOwn page and chunk records'),
 (600,270,960,145,'Recorded relationship','A may reference B without merging their text\nSame feature or number alone does not prove a relationship')
], [((400,105),(600,105)),((1030,105),(1210,105)),((210,180),(210,270)),((400,342),(600,342))])
diagram('architecture.png',(1600,660),[
 (20,230,350,160,'Browser','Questions and filters\nCitations and source reader\nReview and export'),
 (620,230,350,160,'Python server','Routing and retrieval\nContext and claim checks\nValidated answer layout'),
 (1220,25,350,160,'Local storage','Source pages and metadata\nChroma vectors\nSQLite registry and cache'),
 (1220,250,350,160,'Ollama models','Qwen3 embeddings\nQwen3.5 generation\nQwen3.5 support checks'),
 (1220,475,350,150,'CPU reranker','Quantized MiniLM\nQuery and passage scoring')
], [((370,310),(620,310)),((970,280),(1220,105)),((970,310),(1220,330)),((970,345),(1220,550))])

diagram('two-flows.png',(1600,460),[
 (20,25,380,155,'Prepare sources','Circulars and inventories\nPages and reviewed text'),
 (610,25,380,155,'Create representations','Owned chunks and metadata\nEmbeddings and keywords'),
 (1200,25,380,155,'Reusable indexes','Chroma and SQLite\nBM25 in memory'),
 (20,280,380,155,'Receive a question','Route and filters\nQuery embedding'),
 (610,280,380,155,'Select evidence','Search existing indexes\nRerank or select sections'),
 (1200,280,380,155,'Produce the answer','Bounded cited context\nGenerate check and format')
], [((400,103),(610,103)),((990,103),(1200,103)),((400,358),(610,358)),((990,358),(1200,358))])

doc=Document();sec=doc.sections[0]
# Remove decorative paragraph borders inherited from the bundled Word template.
for style in doc.styles:
 for border in list(style.element.iter(qn('w:pBdr'))):
  border.getparent().remove(border)
sec.page_width=Inches(8.2677);sec.page_height=Inches(11.6929)
sec.top_margin=Inches(.67);sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.65)
sec.header_distance=sec.footer_distance=Inches(.27)
width=sec.page_width-sec.left_margin-sec.right_margin
for name,size in [('Normal',10.3),('Title',27),('Subtitle',15),('Heading 1',19),('Heading 2',12.8),('Caption',9)]:
 style=doc.styles[name];style.font.name='Arial';style.font.size=Pt(size);style.font.color.rgb=RGBColor(0,0,0)
 style.paragraph_format.space_after=Pt(7);style.paragraph_format.line_spacing=1.13
 if name.startswith('Heading'):
  style.font.bold=True;style.paragraph_format.space_before=Pt(11);style.paragraph_format.keep_with_next=True
doc.styles['Caption'].font.italic=True
hp=sec.header.paragraphs[0];hp.text='CIRCULAR INTELLIGENCE     CORE PRODUCT GUIDE';hp.style='Normal'
for run in hp.runs:run.font.size=Pt(8);run.font.color.rgb=RGBColor(0,0,0)
fp=sec.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
run=fp.add_run('v0.5.1  |  27 September 2026  |  ');run.font.size=Pt(8)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');fp._p.append(field)
doc.core_properties.title='Circular Intelligence Core Product Guide'
doc.core_properties.subject='Implemented local NPCI circular research product through version 0.5.1'
doc.core_properties.author='Circular Intelligence capstone project'

def inline(p,text):
 for i,piece in enumerate(re.split(r'(\*\*.*?\*\*|`[^`]+`)',text)):
  run=p.add_run(piece[2:-2] if piece.startswith('**') else piece[1:-1] if piece.startswith('`') else piece)
  if piece.startswith('**'):run.bold=True
  if piece.startswith('`'):run.font.name='Consolas';run.font.size=Pt(9)

def column_widths(headers):
 if len(headers)==2:
  if headers[1]=='Page':return [.89,.11]
  if headers[1]=='Count':return [.80,.20]
  if headers[0].startswith('Source in'):return [.52,.48]
  return [.35,.65]
 if headers[0]=='Product membership':return [.50,.25,.25]
 if headers[0]=='Validation layer':return [.24,.34,.42]
 if headers[0]=='Priority':return [.12,.42,.46]
 return [.24,.41,.35]

def table(rows):
 widths=column_widths(rows[0]);t=doc.add_table(rows=1,cols=len(rows[0]));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for col,w in zip(t.columns,widths):col.width=int(width*w)
 for i,row in enumerate(rows):
  cells=t.rows[0].cells if i==0 else t.add_row().cells
  trPr=t.rows[i]._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trPr.append(cant)
  if i==0:
   repeat=OxmlElement('w:tblHeader');trPr.append(repeat)
  for j,(cell,value) in enumerate(zip(cells,row)):
   cell.width=int(width*widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   tcPr=cell._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
   for edge in ('top','left','bottom','right','insideH','insideV'):
    e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
   tcPr.append(borders);margins=OxmlElement('w:tcMar')
   for edge,value_margin in [('top',65),('bottom',65),('left',85),('right',85)]:
    e=OxmlElement('w:'+edge);e.set(qn('w:w'),str(value_margin));e.set(qn('w:type'),'dxa');margins.append(e)
   tcPr.append(margins);shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'33463F' if i==0 else ('F2F4F3' if i%2==0 else 'FFFFFF'));tcPr.append(shade)
   p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.05
   run=p.add_run(value);run.font.size=Pt(9);run.font.bold=i==0
   run.font.color.rgb=RGBColor(255,255,255) if i==0 else RGBColor(0,0,0)
   if (j>0 and re.fullmatch(r'[\d,]+',value)) or rows[0][j] in ('Page','Count','Priority'):p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 doc.add_paragraph().paragraph_format.space_after=Pt(1)

source=(ROOT/'docs/PRODUCT_DOCUMENT.md').read_text(encoding='utf-8-sig')
pages=source.split('<!--PAGE-->');page_titles=[]
for page_index,page in enumerate(pages):
 lines=page.strip().splitlines();i=0;first_heading=True
 while i<len(lines):
  line=lines[i].strip()
  if not line:i+=1;continue
  if line.startswith('# '):
   title=line[2:];page_titles.append(title)
   p=doc.add_paragraph(title,'Title' if page_index==0 else 'Heading 1')
   p.paragraph_format.space_before=Pt(0)
   if page_index:p.paragraph_format.page_break_before=True
   first_heading=False;i+=1;continue
  if line.startswith('## '):
   doc.add_paragraph(line[3:],'Subtitle' if page_index==0 and line[3:]=='Core product guide' else 'Heading 2');i+=1;continue
  if line.startswith('|'):
   rows=[]
   while i<len(lines) and lines[i].strip().startswith('|'):
    cells=[v.strip() for v in lines[i].strip().strip('|').split('|')]
    if not all(re.fullmatch(r':?-+:?',c.replace(' ','')) for c in cells):rows.append(cells)
    i+=1
   table(rows);continue
  if line.startswith('```'):
   i+=1;code=[]
   while i<len(lines) and not lines[i].startswith('```'):code.append(lines[i]);i+=1
   p=doc.add_paragraph();p.paragraph_format.space_after=Pt(8)
   for j,cl in enumerate(code):
    r=p.add_run(cl+('\n' if j<len(code)-1 else ''));r.font.name='Consolas';r.font.size=Pt(9)
   i+=1;continue
  image=re.match(r'!\[(.*?)\]\((.*?)\)',line)
  if image:
   path=ROOT/'docs'/image[2];im=Image.open(path);w,h=im.size
   maxh=4.15 if path.name in ('home.png','answer.png') else 3.25
   imgw=min(6.96,maxh*w/h)
   p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
   p.add_run().add_picture(str(path),width=Inches(imgw));i+=1;continue
  p=doc.add_paragraph();inline(p,line)
  if re.match(r'^\d+\. ',line):
   p.paragraph_format.left_indent=Inches(.16);p.paragraph_format.first_line_indent=Inches(-.16)
  if line.startswith('**Figure '):p.style='Caption'
  i+=1
doc.save(OUT)
(ROOT/'docs/product-assets/page_titles.json').write_text(json.dumps(page_titles,indent=2),encoding='utf-8')
print(str(OUT));print('Planned pages:',len(pages),'Words:',len(re.findall(r'\b\w+\b',source)))

