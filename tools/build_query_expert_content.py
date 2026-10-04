"""Author the concise midterm narrative from a pinned experiment record."""
import json, re
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'output/midsem'; BUILD=ROOT/'tmp/midsem'; stamp='20261003T172349Z'
result=json.loads((OUT/'evidence'/stamp/'retrieval.json').read_text(encoding='utf-8'))
meta=json.loads((OUT/'team_details.json').read_text())
title='Payments Guidelines Query Expert'
team=meta.get('team_id') or '[Team ID]'
members=' / '.join(meta['members']) if len(meta.get('members',[]))==5 and all(meta['members']) else '[Team member names]'
table=[['Retrieval method','Expected evidence found','Negatives with no result','Median search ms']]
for key,name in [('bm25','BM25'),('vector','Vector'),('hybrid','Hybrid'),('hybrid_reranked','Hybrid + reranker')]:
 s=result['runs'][key]['summary'];table.append([name,f"{s['anchors_found']}/{s['anchors_expected']}",f"{s['negative_empty']}/{s['negative_cases']}",str(s['median_ms'])])
roadmap=[
 ('Weeks 1–2: validate answers','Review the existing 50 cases and freeze unseen questions; deliver a reviewed test set and answer-quality scorecard.'),
 ('Weeks 3–5: improve reliability','Repair source and answer failures, shorten summaries and compare retrieval settings; deliver a tested revision and quality/latency comparison.'),
 ('Weeks 6–8: test with users','Observe research tasks and package setup and the final demo; deliver user-task results and a repeatable local demonstration. Another source is an optional later extension.')]
notes=[
"Our project is Payments Guidelines Query Expert. It helps payments professionals ask questions about documented requirements and inspect the sources behind an answer. We have started with NPCI circulars across multiple payment products. The intended product is broader, but this is the coverage implemented today. For this midterm, we are presenting a working local prototype, its design choices, a preliminary retrieval comparison and the work still needed to evaluate answer quality. The product name describes the task it supports. The output remains AI assistance that a user should review, rather than an authoritative compliance decision.",
"Our intended users are product, operations and compliance teams in banks and fintechs. A question about one payment feature may require several circulars, including later clarifications. A user must first find the right documents, then identify the relevant conditions, and finally check which source supports each statement. Manual searching makes this slow, while a fluent AI answer can hide missing conditions. We want to make this question-to-source journey easier. We have not yet measured time savings. A useful outcome would be a user finding the necessary requirement with fewer searches while still checking the original guideline before acting.",
"The prototype supports three tasks: ask a question, inspect its supporting source, and summarise or compare selected circulars. The screenshot is from the actual application. In the recorded example, we ask who checks UPI Tap and Pay enablement before each transaction. The generated answer identifies the Remitter Bank and provides the supporting passage from OC 186A. Clicking the source opens the document reader. This demonstrates traceability for one question, rather than general accuracy. The interface now starts with the question; model settings and collection details are secondary. It also offers source-only search when a generated answer is unnecessary or the model is unavailable.",
"The initial knowledge base contains 1,562 searchable NPCI documents. We ingest Markdown text together with the product inventory. Each passage retains its document and page identity, so text from different circulars is not merged into one anonymous block. Chunking follows page and clause boundaries where recognised; longer blocks use 220-word windows with 35-word overlap. Chroma stores embeddings for semantic search, and SQLite holds metadata and document links. The source register also records 359 unavailable or excluded entries. This matters because an absent source must not look like an absence of requirements. Noisy extraction and uncertain dates remain data limitations.",
"At question time, product and document filters first narrow the search. BM25 finds exact words and circular references, while Qwen embeddings find passages with similar meaning. Hybrid retrieval combines these results. A local MiniLM reranker scores candidates for focused questions. Summary and comparison routes instead select evidence across document sections and bypass reranking. We assemble a bounded set of passages and pass it to Qwen3.5 through Ollama. Citation and support checks run before display, and the reader can inspect the sources. We use pretrained models without fine-tuning because our first requirement is access to attributable documents. Local inference avoids paid model calls, but the 32 GB machine limits speed and context. Automated support checks use the same model and can share its mistakes.",
"We compared BM25, vector, hybrid and hybrid with reranking using the same corpus and a frozen 50-case UPI dataset. All methods found all 55 expected source anchors. This measures evidence retrieval, not answer correctness. BM25 returned no evidence for all six negative cases; the semantic methods returned irrelevant passages for one unrelated question. Reranking increased median retrieval time without improving anchor coverage on this set. Many questions explicitly name circulars, and summary and comparison routes bypass reranking, so the dataset does not yet discriminate well between methods. It is AI-assisted and awaits independent review. We also retained six end-to-end scenarios: four generated responses and two pre-generation abstentions. The summary and comparison were longer and slower than desired, giving us a concrete improvement area.",
"Our midterm milestone is the working local application, a reproducible experiment record, and this presentation and report. The next steps follow one sequence. First, independently review reference answers and score correctness, citation support and missing conditions. Second, use the failures to repair sources and compare retrieval settings on harder cases. Third, observe payments users performing research tasks and pilot one additional guidelines source with its own metadata and evaluation. That expansion is not implemented yet. Reviewer availability and source quality are the main dependencies. We will measure answer quality separately from retrieval coverage, and balance any improvement against response latency. Cloud hosting or fine-tuning is not a prerequisite for these next steps.",
"We would like feedback on three decisions. First, is this focused question-and-source workflow a useful and manageable capstone scope for payments teams? Second, which questions should the independently reviewed evaluation prioritise, especially when requirements span several documents? Third, what evidence would convince you that the assistant is useful in practice: task completion, fewer corrections, time spent, or a combination? Our immediate priority is to validate the answers and improve the failures before expanding coverage. The current prototype demonstrates a complete local RAG flow. The remaining work is to establish how reliably it answers realistic questions and how effectively users can verify those answers."
]
notes[2]="The prototype helps users ask questions, inspect sources, and summarise or compare circulars. Here the user asks: Can I use UPI Tap and Pay at a shop's card machine? No circular number is needed in the question. The recorded local response says the terminal must be NFC-enabled and appropriately certified, and cites the supporting clause. NFC is the contactless technology used for tapping a device. The response also provides merchant-only and integration conditions. The slide uses enlarged captures from the application rendering of this saved response; the answer text has not been rewritten for the presentation. This shows how the user can follow an answer back to a circular. It is one illustrative case, not an accuracy result."
notes[3]="The knowledge base starts with Markdown circulars and product inventory records. We preserve each document and page before splitting its text. A recognised heading or numbered clause begins a new block. A long block is split into windows of at most 220 words, with 35 words repeated between adjacent windows to preserve some context. We never combine text from different circulars into one chunk. Each chunk retains its document ID, page, heading and links to its neighbours. The same text feeds BM25 keyword search and Qwen embeddings stored in Chroma. SQLite stores identities, filters and relationships. This structure lets us search by words or meaning while tracing every selected passage back to its source."
notes[5]="This experiment asks whether more complex retrieval gives us better evidence than a simple keyword baseline. All four methods ran the same 50 UPI cases on the same corpus. Expected evidence found means the required source passages appeared in the retrieved context: all methods found 55 out of 55. It does not mean 55 correct AI answers. The second measure tests six cases that should return no evidence. BM25 returned nothing in all six; each semantic variant returned irrelevant passages in one case. The last column measures retrieval time only, excluding answer generation. BM25 was fastest; adding reranking to hybrid increased the median from 78 to 160.5 milliseconds without improving coverage here. The conclusion is to keep BM25 as a strong baseline and test harder cases before claiming a benefit from extra complexity. Many queries name circulars and 15 coverage routes bypass reranking. The AI-assisted dataset still needs independent review, so this cannot determine overall answer quality or the best method for every question."
notes[6]="The roadmap now identifies work, timing and a concrete output. The midterm milestone is complete at the prototype level: cited questions and answers, source inspection, summaries and comparisons. In the first two weeks after midterm, we plan to review the existing 50 cases with an independent reviewer and freeze unseen questions. The output is a reviewed test set and an answer-quality scorecard. In weeks three to five, we will fix source and answer failures, improve overly long summaries, and compare retrieval choices. The output is a tested revision with a documented quality and latency comparison. In weeks six to eight, payments users will try research tasks, and we will prepare reproducible setup and a final demo. The output is a user-task readout and a repeatable local demonstration. These are indicative planning windows, dependent on reviewer access. A second payments source is an optional extension after validation, not a condition for completing the core capstone."
titles=[title,'Problem Statement and Motivation','Proposed Solution','Data Sources and Tools','How the System Works','What the Preliminary Evaluation Tells Us','Implementation Roadmap','Discussion and Feedback']
notes[5]+=" Some questions need more than one passage, which is why there are 55 expected evidence matches across 50 cases."
times=[45,55,60,60,80,80,60,40];speakers=[1,1,2,2,3,4,5,5]
d=dict(title=title,subtitle='Answers to payment questions, with sources you can check',team=team,members=members,stamp=stamp,table=table,notes=notes,titles=titles,times=times,speakers=speakers,roadmap=roadmap,source_paths=['evidence/'+stamp+'/retrieval.json','evidence/'+stamp+'/live.json','evidence/browser-query-expert/report.json'])
d['source_paths']+=['evidence/simple-demo/answer.json','evidence/simple-demo-friend/answer.json']
(BUILD/'query_expert_content.json').write_text(json.dumps(d,indent=2,ensure_ascii=False),encoding='utf-8')
script=['# Payments Guidelines Query Expert','', 'Midterm narration | Suggested five speakers | Planned duration: 8 minutes','', 'Use the revised deck. Speaker assignments are not contribution claims. Replace team details before recording. Actual duration must be checked in rehearsal.','']
elapsed=0
for i,n in enumerate(notes):
 end=elapsed+times[i];script.extend([f'## Slide {i+1} — {titles[i]}',f'Speaker {speakers[i]} | {elapsed//60}:{elapsed%60:02d}–{end//60}:{end%60:02d}','',n,'']);elapsed=end
script.extend(['## Recording checklist','- Fill in team ID, names and the registered title if required; keep all files consistent.','- Rehearse all speakers with the same deck. Keep the final recording between 5 and 10 minutes.','- Check audio, slide transitions and readability from beginning to end.','- Keep the distinction between retrieval coverage and answer correctness.','- Test submission links from another account if sharing by link.','- Submit the recording, slides and report and retain the receipt.','- Recording, rehearsal and portal submission are still pending.'])
(OUT/'Narration_and_Recording_Checklist.md').write_text('\n'.join(script),encoding='utf-8')
doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.7)
normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(11);normal.paragraph_format.space_after=Pt(6);normal.paragraph_format.line_spacing=1.04
for name,size in [('Title',25),('Heading 1',19),('Heading 2',12)]:
 style=doc.styles[name];style.font.name='Calibri';style.font.size=Pt(size);style.font.color.rgb=RGBColor.from_string('000000' if name=='Title' else '117A78')
for style in doc.styles:
 if style.type==1:
  for border in list(style.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
footer=sec.footer.paragraphs[0];footer.text=title+' | Midterm | ';footer.style=doc.styles['Caption']
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
def p(t):doc.add_paragraph(t)
def h(t):doc.add_heading(t,2)
def tbl(headers,rows):
 t=doc.add_table(rows=1,cols=len(headers));t.style='Light Shading Accent 1'
 for c,v in zip(t.rows[0].cells,headers):c.text=v
 for values in rows:
  for c,v in zip(t.add_row().cells,values):c.text=str(v)
 for row in t.rows:
  row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
  for c in row.cells:
   for para in c.paragraphs:
    para.paragraph_format.space_after=Pt(3);para.paragraph_format.space_before=Pt(2)
    for r in para.runs:r.font.size=Pt(10)
 return t
doc.add_heading(title,0)
p('Midterm Capstone Report | 4 October 2026\n'+team+' | '+members)
h('Problem and motivation')
p('Payments teams must interpret requirements spread across circulars, products and revisions. Finding a relevant document is only the first step: users must identify the applicable conditions and check the evidence behind an answer. Manual searching takes effort, while an AI answer can sound convincing despite using the wrong source or omitting a condition.')
p('Payments Guidelines Query Expert is a local AI assistant for product, operations and compliance professionals in banks and fintechs. It helps them ask a payments question, read a cited answer and inspect its source. Reduced research effort is the intended benefit; user time savings have not yet been measured.')
h('Proposed solution and scope')
tbl(['Implemented capability','User outcome'],[
 ('Ask questions','Retrieve relevant passages and generate answers with citations.'),
 ('Inspect guidelines','Search by product or reference, open source text and available original PDFs, and export notes.'),
 ('Summarise and compare','Read selected documents together while preserving each circular’s attribution.')])
p('Initial corpus: NPCI circulars across multiple payment products. Preliminary evaluation focuses on UPI. Other payment networks, regulators and institution-specific guidelines are potential extensions. This submission covers the RAG assistant only. Automated compliance decisions and cross-functional workflow orchestration are outside its scope.')
h('Data and preprocessing')
p('The imported register contains 1,921 records: 1,562 searchable documents and 359 unavailable or excluded entries. Searchable text spans 4,826 pages and 14,060 chunks. Markdown provides the text; product inventory JSON supplies identity, availability and provenance. This is a saved collection, not a live website feed.')
p('A recognised heading or numbered clause starts a text block within a page. Long blocks use windows of at most 220 words with 35-word overlap. Chunks never mix circulars or cross page boundaries; each retains document ID, page, heading and neighbour links. BM25 indexes chunk text, Qwen embeddings go into Chroma, and SQLite stores metadata and relationships. Related circulars are linked, not merged. A reference alone does not establish supersession. Extracted text and dates can be incomplete.')
doc.add_page_break();doc.add_heading('System design and reasoning',1)
h('Preparation and question flow')
p('Preparation: Markdown and inventory → attributed chunks → keyword index, Chroma vectors and SQLite metadata.\nQuestion flow: question and filters → relevant passages → bounded evidence → local answer generation → citation and support checks → answer with inspectable sources.')
tbl(['Module','Why it is included','Trade off to evaluate'],[
 ('Chunking and metadata','Keep passages focused and preserve document identity, pages and links.','Small chunks can miss conditions; extraction may be incomplete.'),
 ('BM25 keyword search','Find exact terminology and circular references.','May miss a differently worded question.'),
 ('Qwen embeddings and Chroma','Find passages with similar meaning.','Similar passages can still be irrelevant.'),
 ('Hybrid retrieval and MiniLM reranking','Combine keyword and semantic results, then reorder candidates for focused questions.','Adds time; benefit over BM25 is not yet demonstrated.'),
 ('Bounded context assembly','Fit selected passages and related evidence within model limits.','Relevant material may be omitted. Summaries and comparisons use section coverage and bypass reranking.'),
 ('Qwen3.5 and support checks','Draft a readable answer; check citations, numbers and quoted support.','Same-model checks can repeat generation errors.'),
 ('Source reader and export','Let users verify and reuse the answer with attribution.','A citation alone does not establish correctness.')])
h('Models and deployment')
p('The frontend uses HTML, CSS and JavaScript with a Python backend. Ollama runs Qwen3.5 9B for generation and Qwen3 embedding 0.6B for semantic retrieval. MiniLM L6 v2 reranks locally through ONNX. SQLite stores metadata; Chroma stores vectors.')
p('We use pretrained models without fine-tuning because the initial requirement is access to changing, attributable documents. The local machine has about 32 GB RAM. Generation uses a 16,384-token context, with a separate 18,000-character evidence budget and up to 3,072 output tokens. Local execution avoids paid model API calls but still has hardware and latency costs. Multi-user capacity has not been evaluated.')
h('Design decisions and evidence')
p('The reasons above explain implementation choices, not proven superiority. The preliminary comparison on the next page tests retrieval behaviour. Independent answer review and controlled comparisons on harder questions are required before claiming that added complexity improves quality.')
doc.add_page_break();doc.add_heading('Progress evaluation and next steps',1)
h('Implemented progress')
p('The app supports cited answers, source inspection, summaries, comparisons and export. The slide demonstration asks about using Tap & Pay at a shop’s card machine and returns cited terminal conditions. Browser checks cover the core journeys and recovery; six earlier scenarios are also retained. A new question about sending money to a friend produced a general feature overview instead of a direct answer, showing why citation support alone is insufficient. Long summaries and slow comparisons remain improvement areas.')
h('Preliminary retrieval comparison')
p('Purpose: determine whether extra retrieval complexity improves evidence selection over a keyword baseline. Four methods ran the same frozen 50-case UPI set on the same corpus. The AI-assisted set contains 55 expected source passages and six negative cases; independent review is pending.')
tbl(table[0],table[1:])
p('Some questions need multiple passages: 50 cases therefore contain 55 expected evidence matches. All methods found 55/55. The negative-case column checks whether search returns no passages when none are expected: BM25 did so in 6/6 cases; semantic methods returned irrelevant passages in one. Search time is the median retrieval time and excludes answer generation.')
p('Inference: BM25 is a strong baseline on this set. Reranking raised hybrid median search time from 78 to 160.5 ms without improving measured coverage. This is not an answer-accuracy score or proof that BM25 is always best. Many queries name circulars; 15 summary/comparison routes bypass reranking. Timing is one preliminary warmed pass. Harder, independently reviewed questions are needed before selecting a method on quality grounds.')
h('Next steps and completion evidence')
tbl(['Indicative timing','Implementation work','Completion evidence'],[
 ('Weeks 1–2 after midterm','Independently review the 50 cases and freeze unseen questions.','Reviewed test set and answer-quality scorecard: correctness, citations, omissions and abstention.'),
 ('Weeks 3–5','Repair source and answer failures, shorten long summaries, compare retrieval settings.','Tested revision and a documented quality versus latency comparison.'),
 ('Weeks 6–8','Observe payments users on research tasks; prepare repeatable setup and final demo.','Task outcomes, corrections and time; a reproducible local demonstration.')])
p('Reviewer access and source completeness are dependencies. A second payments source is optional after validation, beyond the core completion gate. Feedback requested: evaluation priorities and evidence of user value.')
doc.save(OUT/'Payments_Guidelines_Query_Expert_Report.docx')
print('Report and narrative authored; words in script:',sum(len(re.findall(r"\b[\w’]+\b",n)) for n in notes))
